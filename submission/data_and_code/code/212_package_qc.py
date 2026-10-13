#!/usr/bin/env python3.11
# ============================================================
# 投稿包终检 / Final QC of the three-journal submission package
#
# 检查项 / Checks:
#   1. 英文交付文件无中文残留(中文工作文档与中文摘要除外)
#   2. 无错误邮箱、无旧题目、无 "Extended Data"(GCB)与 "(ref. N)" 残留
#   3. 三刊题目词数(NC <= 15)
#   4. 关键数字三版一致
#   5. GCB 作者-年份引文均能在文献表中找到
#   6. 列出所有须作者补充的占位符
# Usage: python3.11 code/212_package_qc.py
# ============================================================
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document

SUB = Path(__file__).resolve().parent.parent / "analysis_v2" / "submission"
CN_OK = {"README.md", "SUBMISSION_CHECKLIST.md", "THREE_VERSIONS.md", "CRediT_and_authors_TEMPLATE.md",
         "FIGURE_MANIFEST.md", "Translated_Abstract_zh.md", "Translated_Abstract_zh.docx",
         "00_UPLOAD_GUIDE.md"}                                    # 中文上传说明 / Chinese upload guides
MAIN = {k: SUB / k / f"Manuscript_{k}_maintext.md" for k in ("GCB", "NEE", "NC")}
fails = 0


def text_of(p: Path) -> str:
    if p.suffix == ".docx":
        d = Document(str(p))
        return "\n".join(x.text for x in d.paragraphs) + "\n".join(
            c.text for t in d.tables for r in t.rows for c in r.cells)
    return p.read_text(encoding="utf-8", errors="ignore")


def report(ok: bool, msg: str) -> None:
    global fails
    fails += 0 if ok else 1
    print(("  ✓ " if ok else "  ✗ ") + msg)


files = [p for p in SUB.rglob("*") if p.suffix in (".md", ".docx")
         and not ({"presentation", "work", "data_and_code"} & set(p.relative_to(SUB).parts))]
print(f"[1] 中文残留(英文交付文件, 共 {len(files)} 个文件)")
bad = [f"{p.relative_to(SUB)}({len(re.findall(r'[一-鿿]', text_of(p)))})" for p in files
       if p.name not in CN_OK and re.search(r"[一-鿿]", text_of(p))]
report(not bad, "无中文残留" if not bad else "含中文: " + ", ".join(bad))

print("[2] 残留项")
alltxt = {p: text_of(p) for p in files}
for pat, label in [("dingchenchen@pku.edu.cn", "错误邮箱"),
                   ("generation of new bird distribution records once records", "旧题目"),
                   (r"\(ref\. \d", "(ref. N) 引用")]:
    hit = [str(p.relative_to(SUB)) for p, t in alltxt.items() if re.search(pat, t)]
    report(not hit, f"{label}: {'无' if not hit else hit}")
gcb_ed = [str(p.relative_to(SUB)) for p, t in alltxt.items() if "GCB" in p.parts and "Extended Data" in t]
report(not gcb_ed, f"GCB 中 'Extended Data': {'无' if not gcb_ed else gcb_ed}")

print("[2b] Markdown 标题前须有空行(否则 pandoc 渲染为正文, Word 中标题消失)")
glued = [str(p.relative_to(SUB)) for p in files if p.suffix == ".md"
         and __import__("re").search(r"[^\n]\n#{1,6} ", p.read_text(encoding="utf-8"))]
report(not glued, "全部标题前有空行" if not glued else f"标题紧贴上段: {glued}")

print("[3] 题目词数")
for k, f in MAIN.items():
    title = re.search(r"^# (.+)$", f.read_text(encoding="utf-8"), re.M).group(1)
    n = len(title.split())
    report(n <= 15 or k != "NC", f"{k}: {n} 词")

print("[4] 关键数字三版一致")
KEYS = ["1,029", "175,901", "649", "392", "1.36", "1.40", "0.85", "1.157", "0.980", "2.23", "2.18", "1.07", "37%", "2.09"]
txt = {k: f.read_text(encoding="utf-8") for k, f in MAIN.items()}
missing = {k: [x for x in KEYS if x not in t] for k, t in txt.items()}
report(all(not v for v in missing.values()),
       "全部出现于三版" if all(not v for v in missing.values()) else f"缺失: {missing}")

print("[5] GCB 作者-年份引文完整性")
g = txt["GCB"]; body, _, refs = g.partition("## References")
entries = [e.strip() for e in refs.split("\n\n") if e.strip()]
keys = set()
for e in entries:
    sur = re.match(r"([A-Z][A-Za-zÀ-ɏ'\-]+)", e)
    yrs = re.findall(r"\((\d{4})[a-z]?\)", e) or re.findall(r"\b(19\d{2}|20\d{2})\b", e)
    if sur and yrs:
        keys.add((sur.group(1), yrs[-1]))
cited = set()
for c in re.findall(r"\(([^()]*?(?:19|20)\d{2}[a-z]?[^()]*?)\)", body):
    for sur, yr in re.findall(r"([A-Z][A-Za-zÀ-ɏ'\-]+)(?:\s+(?:et al\.|&\s*[A-Z][A-Za-z'\-]+))?,?\s*((?:19|20)\d{2})", c):
        cited.add((sur, yr))
corporate = {("World", "2024"), ("Platform", "2024"), ("International", "2024")}   # 机构作者的尾词
unresolved = sorted(c for c in cited - keys - corporate
                    if c[0] not in {"CI", "SSP", "CMIP", "SD", "AIC", "China", "Supplementary", "Figure", "Table"})
report(not unresolved, f"正文引文均有对应条目" if not unresolved else f"无对应条目: {unresolved}")

print("[5b] NEE/NC 数字引文: 首次出现顺序、全部被引、无非数字上标(如 [A1] 占位)")


def _expand(tok: str) -> list[int]:
    out: list[int] = []
    for p in tok.split(","):
        p = p.strip()
        if re.fullmatch(r"\d+", p):
            out.append(int(p))
        elif re.fullmatch(r"\d+[–-]\d+", p):
            a, b = map(int, re.split(r"[–-]", p))
            out += list(range(a, b + 1))
    return out


for k in ("NEE", "NC"):
    _, _, rest = txt[k].partition("## Abstract")          # 跳过作者单位上标 / skip the affiliation marker
    body, _, refs = rest.partition("## References")
    refs, _, tail = refs.partition("\n## ")
    n_ref = len(re.findall(r"^\d+\.\s", refs, re.M))
    seq = [x for tok in re.findall(r"\^([\d,–\-]+)\^", body + tail) for x in _expand(tok)]
    first = list(dict.fromkeys(seq))
    odd = [s for s in re.findall(r"\^([^\^\n]{1,20})\^", body + tail) if not re.fullmatch(r"[\d,–\-]+", s)]
    ok = first == sorted(first) and set(seq) == set(range(1, n_ref + 1)) and not odd
    report(ok, f"{k}: {n_ref} 条, 被引 {len(set(seq))}, 首现有序 {first == sorted(first)}, 非数字上标 {odd or '无'}")

print("[5c] 已撤回的结论不得再现 / retired claims must not reappear")
# 1,056 / 1,059 是 CBNR 预发布版计数; 对外一律用 v2 的 1,029
# NB: 1,056/1,059 are pre-release CBNR counts; the published v2 has 1,029 records
RETIRED = [r"strongly biased toward the east", r"sharpest available evidence", r"\b1,05[69]\b"]
# 根目录 Manuscript_NatureFamily_* 是历史源稿(非投稿件, README 已注明), 不参与此项
# NB: the root Manuscript_NatureFamily_* files are the historical source, not submission items
# 内部说明文档会记录更正过程(如「1,056 → 1,029」), 不属于投稿件, 不参与此项
# NB: internal notes document the corrections themselves and are not submission items
DOC_ONLY = {"README.md", "SUBMISSION_CHECKLIST.md", "THREE_VERSIONS.md", "FIGURE_MANIFEST.md"}
hit = [f"{p.relative_to(SUB)}: {pat}" for p, t in alltxt.items() for pat in RETIRED
       if re.search(pat, t) and not p.name.startswith("Manuscript_NatureFamily_") and p.name not in DOC_ONLY]
report(not hit, "无" if not hit else f"残留: {hit}")

print("[7] 上传件完整性 / upload set completeness (215 的产物)")
import zipfile  # noqa: E402

REQUIRED = {
    "GCB": ["00_UPLOAD_GUIDE.md", "Manuscript_GCB_main_document.docx", "pdf/Supporting_Information_GCB.pdf",
            "cover_letter_GCB.docx", "Translated_Abstract_zh.docx", "pdf/GCB_manuscript_with_figures.pdf"]
           + [f"figures/Figure{i}.pdf" for i in range(1, 8)],
    "NEE": ["00_UPLOAD_GUIDE.md", "Manuscript_NEE_complete.docx", "pdf/Supplementary_Information_NEE.pdf",
            "cover_letter_NEE.docx", "ReportingSummary_DRAFT_NEE.docx", "pdf/NEE_manuscript_with_figures.pdf"]
           + [f"figures/Fig{i}.pdf" for i in range(1, 6)],
    "NC": ["00_UPLOAD_GUIDE.md", "Manuscript_NC_complete.docx", "pdf/Supplementary_Information_NC.pdf",
           "cover_letter_NC.docx", "ReportingSummary_DRAFT_NC.docx", "pdf/NC_manuscript_with_figures.pdf"]
          + [f"figures/Fig{i}.pdf" for i in range(1, 8)],
}
for k, rel in REQUIRED.items():
    miss = [r for r in rel if not (SUB / k / r).exists() or (SUB / k / r).stat().st_size == 0]
    zips = sorted((SUB / "_archives").glob(f"{k}_upload_*.zip"))
    zbad = not zips or zipfile.ZipFile(zips[-1]).testzip() is not None
    report(not miss and not zbad, f"{k}: {len(rel) - len(miss)}/{len(rel)} 件齐全"
           + (f", 缺 {miss}" if miss else "") + (f"; 压缩包 {zips[-1].name if zips else '缺失'}" + (" 损坏" if zbad and zips else "")))
pres = SUB / "presentation"
decks = [pres / "AOC2026_Singapore_EN/AOC2026_ChenchenDing_BirdsOrBirders_20slides.pptx",
         pres / "AOC2026_Singapore_EN/AOC2026_speaker_script_EN-CN.docx"] + sorted(pres.glob("鸟类新纪录生成风险_会议报告_20页_v*.pptx"))
old = [p.name for p in pres.glob("鸟类新纪录生成风险_会议报告_20页_v*.pptx")][:-1] if len(list(pres.glob("鸟类新纪录生成风险_会议报告_20页_v*.pptx"))) > 1 else []
report(all(p.exists() for p in decks) and len(decks) >= 3 and not old,
       f"会议报告: {', '.join(p.name for p in decks)}" + (f";多版本并存 {old}" if old else ""))

print("[8] 可编辑图件 / editable figures (217)")
import csv  # noqa: E402

SI_N = {"GCB": 21, "NEE": 23, "NC": 20}
man = SUB / "editable_figures_manifest.csv"
mrows = list(csv.DictReader(man.open(encoding="utf-8"))) if man.exists() else []
for k in ("GCB", "NEE", "NC"):
    pdfs = sorted((SUB / k / "figures").glob("*.pdf"))
    miss = [f.stem for f in pdfs if not f.with_suffix(".pptx").exists()]
    n_si = len(list((SUB / k / "figures_SI").glob("FigS*.pptx")))
    report(not miss and n_si == SI_N[k], f"{k}: 主图 {len(pdfs) - len(miss)}/{len(pdfs)} 有 PPTX, SI {n_si}/{SI_N[k]}"
           + (f", 缺 {miss}" if miss else ""))
bad = [f"{r['journal']} {r['label']}" for r in mrows
       if float(r["mean_abs_diff"]) >= 12 or float(r["pct_px_diff_gt64"]) >= 8]
report(bool(mrows) and not bad, f"回渲染差异均在阈值内({len(mrows)} 个)" if not bad else f"需目检: {bad}")

print("[9] 顶刊版主图 / polished figures")
for k in ("GCB", "NEE", "NC"):
    names = [f.stem for f in sorted((SUB / k / "figures").glob("*.pdf"))]
    pol = SUB / k / "figures_polished"
    miss = [n for n in names for e in ("pdf", "png", "pptx") if not (pol / f"{n}.{e}").exists()]
    report(not miss, f"{k}: {len(names)} 张主图的 PDF/PNG/PPTX " + ("齐全" if not miss else f"缺 {sorted(set(miss))}"))

print("[10] 数据与代码 / data and code (219)")
dc = SUB / "data_and_code"
need = ["README.md", "LICENSE", "config.R", "setup_workspace.R", "tests/smoke_test.R", "environment/renv.lock",
        "environment/R_sessionInfo.txt"]
miss = [n for n in need if not (dc / n).exists()]
n_code = len(list((dc / "code").glob("*"))) if (dc / "code").exists() else 0
n_data = len([f for f in (dc / "data").rglob("*") if f.is_file()]) if (dc / "data").exists() else 0
sd_ok = all(len(list((dc / "source_data" / k).glob("*.csv"))) == len(list((SUB / k / "figures").glob("*.pdf")))
            for k in ("GCB", "NEE", "NC"))
zips_dc = sorted((SUB / "_archives").glob("data_and_code_*.zip"))
zip_ok = bool(zips_dc) and zipfile.ZipFile(zips_dc[-1]).testzip() is None
report(not miss and n_code > 50 and n_data > 20 and sd_ok and zip_ok,
       f"代码 {n_code}、数据 {n_data}、源数据与主图一一对应 {sd_ok}、压缩包 {zips_dc[-1].name if zips_dc else '缺'}"
       + (f";缺 {miss}" if miss else ""))

print("[11] 新版报告 / v3 talk decks")
v3 = SUB / "presentation" / "AOC2026_v3_redesign"
n_pp, n_pdf = len(list(v3.glob("*.pptx"))), len(list(v3.glob("*.pdf")))
n_nt = len(list(v3.glob("AOC2026_v3_speaker_notes_*.md")))
report(n_pp == 4 and n_pdf == 4 and n_nt == 2, f"PPTX {n_pp}/4、PDF {n_pdf}/4、讲稿 {n_nt}/2;旧版保留: "
       + ", ".join(p.name for p in sorted((SUB / "presentation").rglob("*.pptx")) if "v3" not in p.name))

print("[6] 须作者补充的占位符")
ph = {}
for p, t in alltxt.items():
    for m in re.findall(r"\[(?:AUTHORS|Authors|To be|Funding|Optional)[^\]]*\]", t):
        ph.setdefault(m[:90], set()).add(str(p.relative_to(SUB)).split("/")[0])
for m, where in sorted(ph.items()):
    print(f"     · {m}  ←  {', '.join(sorted(where))}")

print(f"\n合计失败项: {fails}")
sys.exit(1 if fails else 0)
