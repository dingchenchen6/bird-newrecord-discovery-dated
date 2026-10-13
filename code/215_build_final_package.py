#!/usr/bin/env python3.11
# ============================================================
# 投稿包终装 / Assemble the final, self-contained submission package
#
# 科学问题之外的交付问题 / Delivery question:
#   三刊投稿包要做到「打开即可上传」: 每刊一份完整稿件(正文 + 表 + 图注, Nature 系嵌图)、
#   支持信息、投稿信、必需表格、预览 PDF、上传说明; 会议 PPT 与讲稿与正文口径一致。
#   Make each journal folder upload-ready and keep the decks in step with the manuscripts.
#
# 输入 / Input:
#   analysis_v2/submission/{GCB,NEE,NC}/(正文、表与图注、SI、投稿信的 md/docx 与 figures/)
#   analysis_v2/submission/ReportingSummary_DRAFT_NatureFamily.{md,docx}
#   ~/.claude/skills/easyslides/projects/{bird_talk_aoc2026_en_*,bird_talk_v2_*}/exports(两份 deck)
# 流程 / Workflow:
#   1. 完整稿件: 正文 + 表 + 图注(NEE/NC 图嵌在图注前; GCB 图单独上传)  -> *_complete / *_main_document
#   2. 投稿信 md -> docx(Times New Roman 12 pt, 单倍行距)
#   3. 预览 PDF(LibreOffice): 完整稿件(含图)、SI、投稿信              -> {J}/pdf/
#   4. Nature 系把 Reporting Summary 草稿复制进期刊目录
#   5. 会议 PPT 与讲稿从 deck 工程同步; 旧版移出投稿包归档
#   6. 每刊写 00_UPLOAD_GUIDE.md; 打包 _archives/{J}_upload_{DATE}.zip 与 presentation zip
#      (含可编辑图件、顶刊版主图与 v3 新版报告; --zips-only 只重打压缩包)
#   7. 自检: zip 完整性、PDF 页数、完整稿件的嵌图数
# 输出 / Output: 见上; 压缩包内文件按上传类别编号
# 关键假设 / Key assumptions:
#   - GCB 采用 Wiley 常规结构: 表与图注置于参考文献之后, 图件单独上传
#   - NEE/NC 初投接受单文件 Word/PDF(2026-09-27 核对 nature.com 官方页面)
#   - md 为唯一可编辑源; docx/pdf/zip 均可由本脚本重建
# 运行 / Run (任务根目录 / from task root): python3.11 code/215_build_final_package.py
# ============================================================
from __future__ import annotations

import datetime
import importlib.util
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

V2 = Path(__file__).resolve().parent.parent
SUB = V2 / "analysis_v2" / "submission"
ARCH = SUB / "_archives"
SOFFICE = "/opt/homebrew/bin/soffice"
DATE = datetime.date.today().strftime("%Y%m%d")
DECKS = Path.home() / ".claude/skills/easyslides/projects"
DECK_EN = DECKS / "bird_talk_aoc2026_en_ppt169_20260923/exports"
DECK_ZH = DECKS / "bird_talk_v2_ppt169_20260816/exports"
EN_FILES = ["AOC2026_ChenchenDing_BirdsOrBirders_20slides.pptx", "AOC2026_ChenchenDing_BirdsOrBirders_20slides.pdf",
            "AOC2026_speaker_script_EN-CN.docx", "AOC2026_speaker_script_EN-CN.md"]
ZH_STEM = "鸟类新纪录生成风险_会议报告_20页_v8"
# 2026-10 新版报告: 同一叙事三种风格 + 中文版, 与旧版并存 / v3 talk: three styles + Chinese, kept alongside v2
DECK_V3 = DECKS / "bird_talk_aoc2026_v3_20261011" / "variants"
V3_VARIANTS = ["light_en", "dark_en", "warm_en", "light_zh"]
V3_DIR = "AOC2026_v3_redesign"
FIG_W = "6.3in"

# 复用 205 的渲染与版式(双倍行距、行号、页码)/ reuse 205's renderer and manuscript styling
_spec = importlib.util.spec_from_file_location("r205", V2 / "code" / "205_prune_references.py")
r205 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(r205)
# 复用 206 的三线表样式(本脚本生成的新 docx 在 206 之后产生)/ reuse 206 booktabs for docx created here
_spec6 = importlib.util.spec_from_file_location("r206", V2 / "code" / "206_restyle_tables.py")
r206 = importlib.util.module_from_spec(_spec6)
_spec6.loader.exec_module(r206)

JOURNALS = {
    "GCB": {"fig": "Figure", "fig_label": r"\*\*Figure (\d+) \|", "embed": False,
            "si": "Supporting_Information_GCB", "complete": "Manuscript_GCB_main_document"},
    "NEE": {"fig": "Fig", "fig_label": r"\*\*Fig\. (\d+) \|", "embed": True,
            "si": "Supplementary_Information_NEE", "complete": "Manuscript_NEE_complete"},
    "NC":  {"fig": "Fig", "fig_label": r"\*\*Fig\. (\d+) \|", "embed": True,
            "si": "Supplementary_Information_NC", "complete": "Manuscript_NC_complete"},
}


def log(*a: object) -> None:
    print("[215]", *a, flush=True)


def with_figures(legends_md: str, J: str) -> str:
    """在每条图注前嵌入对应图件 / insert each figure above its legend."""
    cfg, figdir = JOURNALS[J], SUB / J / "figures"

    def add(m: re.Match) -> str:
        png = figdir / f"{cfg['fig']}{m.group(1)}.png"
        if not png.exists():
            sys.exit(f"缺图 / missing figure: {png}")
        return f"![]({png.as_posix()}){{width={FIG_W}}}\n\n{m.group(0)}"
    return re.sub(r"(?m)^" + cfg["fig_label"], add, legends_md)


def complete_md(J: str, embed: bool) -> str:
    """正文 + 表 + 图注 / main text followed by tables and figure legends."""
    main = (SUB / J / f"Manuscript_{J}_maintext.md").read_text(encoding="utf-8").rstrip()
    tl = (SUB / J / f"Manuscript_{J}_tables_and_legends.md").read_text(encoding="utf-8").strip()
    if embed:
        tl = with_figures(tl, J)
    return main + "\n\n" + tl + "\n"


def render_md(md_text: str, out_docx: Path) -> None:
    """渲染并套用稿件版式与三线表 / render with manuscript styling and booktabs tables."""
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td) / (out_docx.stem + ".md")
        tmp.write_text(md_text, encoding="utf-8")
        r205.render(tmp, out_docx)
    doc = Document(str(out_docx))
    for t in doc.tables:
        r206.booktabs(t)
    doc.save(str(out_docx))


def style_letter(path: Path) -> None:
    """投稿信版式: Times New Roman 12 pt, 单倍行距, 段后 6 pt / letter layout."""
    doc = Document(str(path))
    r205.force_fonts(doc, "Times New Roman")               # 去掉主题字体 / drop theme fonts
    for st in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if st in [s.name for s in doc.styles]:
            s = doc.styles[st]
            s.font.name, s.font.size = "Times New Roman", Pt(12)
            if s.element.rPr is not None and s.element.rPr.rFonts is not None:
                s.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            s.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            s.paragraph_format.space_after = Pt(6)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    doc.save(str(path))


def render_letter(md: Path, out: Path) -> None:
    r = subprocess.run(["pandoc", str(md), "-o", str(out), "--from", "markdown", "--to", "docx"],
                       capture_output=True, text=True)
    if r.returncode != 0:
        sys.exit(f"pandoc 失败 / failed ({md.name}): {r.stderr[:300]}")
    style_letter(out)


def to_pdf(docx: Path, outdir: Path, name: str | None = None) -> Path:
    outdir.mkdir(exist_ok=True)
    r = subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(docx)],
                       capture_output=True, text=True, timeout=900)
    pdf = outdir / (docx.stem + ".pdf")
    if r.returncode != 0 or not pdf.exists():
        sys.exit(f"soffice 失败 / failed ({docx.name}): {r.stderr[-300:]}")
    if name:
        final = outdir / name
        pdf.replace(final)
        pdf = final
    return pdf


# ---------------- 上传说明 / upload guides ----------------
GUIDE_HEAD = ("# {J} 上传清单 / Upload guide\n\n"
              "本目录可直接用于投稿。下表「上传类别」为投稿系统中的文件类型;"
              "压缩包 `_archives/{J}_upload_{DATE}.zip` 内的文件按此顺序编号。\n\n"
              "| 编号 | 文件 | 上传类别 | 说明 |\n|---|---|---|---|\n")
GUIDES = {
    "GCB": [
        ("01", "Manuscript_GCB_main_document.docx", "Main Document", "题名页、摘要、关键词、正文、致谢、声明、参考文献,其后为表与图注(Wiley 常规结构)"),
        ("02", "figures/Figure1–7.pdf", "Figure", "矢量 PDF,每图一个文件;450 dpi PNG 为备用"),
        ("03", "pdf/Supporting_Information_GCB.pdf", "Supporting Information", "Notes S1–S7、Tables S1–S10、Figs S1–S21;每张图与图注同页。可编辑版 `Supporting_Information_GCB.docx`"),
        ("04", "cover_letter_GCB.docx", "Cover Letter", "GCB 七问格式;第 7 题推荐审稿人待作者填写"),
        ("05", "Translated_Abstract_zh.docx", "Supplementary Material Not for Review", "中文摘要(GCB 鼓励提供)"),
        ("—", "figures/*.pptx · figures_SI/*.pptx", "按需", "全部主图与 SI 图的可编辑 PPTX(原生形状、文字可改;4 张县级多图地图为嵌图版)"),
        ("—", "figures_polished/", "备选", "顶刊设计版主图(同一数据, 统一配色与版式), PDF + PNG + 可编辑 PPTX"),
        ("—", "../data_and_code/", "按编辑部要求", "数据、代码、各图源数据与运行环境;压缩包 `_archives/data_and_code_日期.zip`"),
        ("—", "pdf/", "不上传", "合作者审阅用 PDF:含图全文、SI、投稿信"),
    ],
    "NEE": [
        ("01", "Manuscript_NEE_complete.docx", "Manuscript", "单文件:正文、Methods、参考文献、声明,其后为表与图注,图嵌在图注前(NEE 初投接受 Word/PDF)"),
        ("02", "pdf/Supplementary_Information_NEE.pdf", "Supplementary Information", "Nature 系要求单一合并 PDF,每张图与图注同页;Notes S1–S11、Tables S1–S12、Figs S1–S23。可编辑版 `Supplementary_Information_NEE.docx`"),
        ("03", "cover_letter_NEE.docx", "Cover Letter", "推荐/回避审稿人待作者补充"),
        ("04", "ReportingSummary_DRAFT_NEE.docx", "Reporting Summary", "逐题答案草稿,须粘贴进官方 Reporting Summary 表格"),
        ("05", "figures/Fig1–5.pdf", "Figure(接收后/修回时)", "矢量 PDF 与 450 dpi PNG;初投已嵌入完整稿件"),
        ("—", "figures/*.pptx · figures_SI/*.pptx", "按需", "全部主图与 SI 图的可编辑 PPTX(原生形状、文字可改;4 张县级多图地图为嵌图版)"),
        ("—", "figures_polished/", "备选", "顶刊设计版主图(同一数据, 统一配色与版式), PDF + PNG + 可编辑 PPTX"),
        ("—", "../data_and_code/", "按编辑部要求", "数据、代码、各图源数据与运行环境;压缩包 `_archives/data_and_code_日期.zip`"),
        ("—", "pdf/", "不上传", "合作者审阅用 PDF"),
    ],
    "NC": [
        ("01", "Manuscript_NC_complete.docx", "Manuscript", "单文件:正文、Methods、参考文献、声明,其后为表与图注,图嵌在图注前(NC 初投格式灵活)"),
        ("02", "pdf/Supplementary_Information_NC.pdf", "Supplementary Information", "Nature 系要求单一合并 PDF,每张图与图注同页;Notes S1–S8、Tables S1–S10、Figs S1–S20。可编辑版 `Supplementary_Information_NC.docx`"),
        ("03", "cover_letter_NC.docx", "Cover Letter", "推荐/回避审稿人待作者补充"),
        ("04", "ReportingSummary_DRAFT_NC.docx", "Reporting Summary", "逐题答案草稿,须粘贴进官方表格"),
        ("05", "figures/Fig1–7.pdf", "Figure(接收后/修回时)", "矢量 PDF 与 450 dpi PNG;初投已嵌入完整稿件"),
        ("—", "figures/*.pptx · figures_SI/*.pptx", "按需", "全部主图与 SI 图的可编辑 PPTX(原生形状、文字可改;4 张县级多图地图为嵌图版)"),
        ("—", "figures_polished/", "备选", "顶刊设计版主图(同一数据, 统一配色与版式), PDF + PNG + 可编辑 PPTX"),
        ("—", "../data_and_code/", "按编辑部要求", "数据、代码、各图源数据与运行环境;压缩包 `_archives/data_and_code_日期.zip`"),
        ("—", "pdf/", "不上传", "合作者审阅用 PDF"),
    ],
}
GUIDE_TAIL = ("\n**编辑与重建**:`*.md` 是唯一可编辑源;改动后依次运行 205(渲染)→ 207(词数)→ 212(终检)"
              "→ 215(本包)。直接在 docx 中修改的内容须回写 md,否则会在下次重建时丢失。\n\n"
              "**投稿前仍须作者补齐**:作者名单与单位、ORCID、CRediT(`../CRediT_and_authors_TEMPLATE.md`)、"
              "基金与项目编号、预印本声明、推荐审稿人;详见 `../SUBMISSION_CHECKLIST.md`。\n")


def write_guide(J: str) -> Path:
    rows = "".join(f"| {a} | `{b}` | {c} | {d} |\n" for a, b, c, d in GUIDES[J])
    p = SUB / J / "00_UPLOAD_GUIDE.md"
    p.write_text(GUIDE_HEAD.format(J=J, DATE=DATE) + rows + GUIDE_TAIL, encoding="utf-8")
    return p


# ---------------- 压缩包 / archives ----------------
def upload_members(J: str) -> list[tuple[Path, str]]:
    d, cfg = SUB / J, JOURNALS[J]
    m: list[tuple[Path, str]] = [(d / "00_UPLOAD_GUIDE.md", "00_UPLOAD_GUIDE.md"),
                                 (d / f"{cfg['complete']}.docx", f"01_{cfg['complete']}.docx")]
    if J == "GCB":
        m += [(f, f"02_Figures/{f.name}") for f in sorted((d / "figures").glob("*.pdf"))]
        m += [(d / "pdf" / f"{cfg['si']}.pdf", f"03_{cfg['si']}.pdf"), (d / "cover_letter_GCB.docx", "04_cover_letter_GCB.docx"),
              (d / "Translated_Abstract_zh.docx", "05_Translated_Abstract_zh.docx")]
    else:
        m += [(d / "pdf" / f"{cfg['si']}.pdf", f"02_{cfg['si']}.pdf"), (d / f"cover_letter_{J}.docx", f"03_cover_letter_{J}.docx"),
              (d / f"ReportingSummary_DRAFT_{J}.docx", f"04_ReportingSummary_DRAFT_{J}.docx")]
        m += [(f, f"05_Figures/{f.name}") for f in sorted((d / "figures").glob("*.pdf"))]
    m += [(f, f"figures_png/{f.name}") for f in sorted((d / "figures").glob("*.png"))]
    # 可编辑图件与顶刊版主图(非上传件, 随包备用)/ editable and polished figures, not for upload
    m += [(f, f"editable_figures/{f.name}") for f in sorted((d / "figures").glob("*.pptx"))]
    m += [(f, f"editable_figures_SI/{f.name}") for f in sorted((d / "figures_SI").glob("*.pptx"))]
    m += [(f, f"figures_polished/{f.name}") for f in sorted((d / "figures_polished").glob("*"))
          if f.suffix in {".pdf", ".png", ".pptx"}]
    m += [(f, f"review_pdf/{f.name}") for f in sorted((d / "pdf").glob("*.pdf")) if cfg["si"] not in f.name]
    return m


def make_zip(path: Path, members: list[tuple[Path, str]]) -> None:
    missing = [str(s) for s, _ in members if not s.exists()]
    if missing:
        sys.exit(f"压缩包缺文件 / missing: {missing}")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for src, arc in members:
            z.write(src, arc)
    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
    if bad:
        sys.exit(f"压缩包损坏 / corrupt member: {bad}")
    log(f"{path.name}: {len(members)} files, {path.stat().st_size / 1e6:.1f} MB, 完整性通过 / integrity OK")


def sync_v3(pres: Path) -> None:
    """新版报告(三风格 + 中文)与讲稿 / v3 decks in three styles plus Chinese, with notes."""
    dst = pres / V3_DIR
    dst.mkdir(exist_ok=True)
    for v in V3_VARIANTS:
        theme, lang = v.split("_")
        stem = f"AOC2026_ChenchenDing_v3_{theme}_{lang.upper()}"
        for ext in ("pptx", "pdf"):
            src = DECK_V3 / v / "exports" / f"{stem}.{ext}"
            if not src.exists():
                sys.exit(f"缺新版报告 / missing v3 deck: {src}")
            shutil.copy2(src, dst / src.name)
    for lang in ("en", "zh"):
        shutil.copy2(DECK_V3 / f"light_{lang}" / "notes" / "total.md", dst / f"AOC2026_v3_speaker_notes_{lang.upper()}.md")
    log(f"新版报告 / v3 decks -> {dst.relative_to(SUB)}: {len(list(dst.glob('*.pptx')))} pptx")


def archives(pres: Path, en_dir: Path) -> None:
    """上传说明与压缩包 / guides and archives."""
    ARCH.mkdir(exist_ok=True)
    for J in JOURNALS:
        write_guide(J)
        for old in ARCH.glob(f"{J}_upload_*.zip"):
            old.unlink()
        make_zip(ARCH / f"{J}_upload_{DATE}.zip", upload_members(J))
    for old in ARCH.glob("presentation_*.zip"):
        old.unlink()
    v3 = pres / V3_DIR
    make_zip(ARCH / f"presentation_{DATE}.zip",
             [(en_dir / f, f"AOC2026_Singapore_EN/{f}") for f in EN_FILES]
             + [(pres / f"{ZH_STEM}.{e}", f"{ZH_STEM}.{e}") for e in ("pptx", "pdf")]
             + [(f, f"{V3_DIR}/{f.name}") for f in sorted(v3.glob("*")) if f.is_file()])


def main() -> None:
    pres, en_dir = SUB / "presentation", SUB / "presentation" / "AOC2026_Singapore_EN"
    if "--zips-only" in sys.argv:          # wm-clean 之后只重打压缩包 / rezip after cleaning
        archives(pres, en_dir)
        log("完成 / done (zips only)")
        return
    # 1–3 稿件、投稿信、PDF / manuscripts, letters, PDFs
    for J, cfg in JOURNALS.items():
        d, pdfdir = SUB / J, SUB / J / "pdf"
        pdfdir.mkdir(exist_ok=True)
        out = d / f"{cfg['complete']}.docx"
        render_md(complete_md(J, cfg["embed"]), out)
        n_img = len(Document(str(out)).inline_shapes)
        log(f"{J}: {out.name}(嵌图 {n_img} 张 / {n_img} embedded figures)")
        if cfg["embed"]:
            review = to_pdf(out, pdfdir, f"{J}_manuscript_with_figures.pdf")
        else:                                            # GCB: 另建含图审阅副本 / GCB review copy with figures
            with tempfile.TemporaryDirectory() as td:
                tmp_docx = Path(td) / f"{J}_manuscript_with_figures.docx"
                with_fig = complete_md(J, True)
                render_md(with_fig, tmp_docx)
                review = to_pdf(tmp_docx, pdfdir)
        render_letter(d / f"cover_letter_{J}.md", d / f"cover_letter_{J}.docx")
        to_pdf(d / f"cover_letter_{J}.docx", pdfdir)
        to_pdf(d / f"{cfg['si']}.docx", pdfdir)
        log(f"{J}: 预览 PDF / review PDFs -> {', '.join(p.name for p in sorted(pdfdir.glob('*.pdf')))}"
            f" | 含图全文 {fitz.open(review).page_count} 页")
    # 4 Reporting Summary
    for J in ("NEE", "NC"):
        for ext in ("docx", "md"):
            shutil.copy2(SUB / f"ReportingSummary_DRAFT_NatureFamily.{ext}", SUB / J / f"ReportingSummary_DRAFT_{J}.{ext}")
    # 5 演示文稿 / presentations
    for f in EN_FILES:
        shutil.copy2(DECK_EN / f, en_dir / f)
    for ext in ("pptx", "pdf"):
        shutil.copy2(DECK_ZH / f"{ZH_STEM}.{ext}", pres / f"{ZH_STEM}.{ext}")
    sync_v3(pres)
    old = sorted(p for p in pres.glob("鸟类新纪录生成风险_会议报告_20页_v*.*") if not p.stem.endswith(ZH_STEM[-2:]))
    if old:
        keep = DECK_ZH / "archive_from_submission"; keep.mkdir(exist_ok=True)
        for p in old:
            shutil.move(str(p), keep / p.name)
        log(f"旧版 deck 移出投稿包 / older decks archived: {[p.name for p in old]} -> {keep}")
    # 6 上传说明与压缩包 / guides and archives
    archives(pres, en_dir)
    log("完成 / done")


if __name__ == "__main__":
    main()
