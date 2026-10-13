#!/usr/bin/env python3.11
# ============================================================
# 按实际引用裁剪并重编号三版文献表 / Prune and renumber the reference lists.
#
# 问题 / Problem:
#   三版的文献表是从长稿整体继承的,而正文分别被裁掉 36–54%。
#   结果:82 条文献里约 55 条从未被引用。期刊要求文献表与引文一一对应,
#   这一条编辑部初审即会发现。NEE/NC 另有一处过时脚注,声称正文存在
#   [A1]–[A6] 上标 —— 那些标签早已被转换成真实编号。
#
# 做法 / Approach:
#   1. 引用集合取「正文 ∪ SI」的并集(SI 无独立文献表,共用主表编号);
#   2. NEE/NC 按**首次出现顺序**重编号 —— 这同时修正了 Nature 系
#      「引用序号须按首次出现排序」的体例要求(裁剪后原序号已不连续);
#   3. GCB 为作者-年份制,只裁剪、不重编号(文献表按姓氏排序);
#   4. 删除过时的 [A1]–[A6] 脚注;
#   5. 由裁剪后的 .md 重新渲染 .docx(排版与脚本 202/203/204 一致)。
#
# Input  / 输入: submission/{J}/Manuscript_{J}_maintext.md 与对应 SI .md
# Output / 输出: 覆盖同名 .md 与 .docx;并打印裁剪前后与验收结果
# Run: python3.11 code/205_prune_references.py
# ============================================================
from __future__ import annotations
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

V2 = Path("/Users/dingchenchen/Documents/New records/bird-new-distribution-records/"
          "tasks/bird_hazard_model_effort_upgrade_v2")
SUB = V2 / "analysis_v2/submission"
BODY_FONT, BODY_PT = "Times New Roman", 12
SUP = re.compile(r"\^([\d,–\-]+)\^")
STALE = re.compile(r"\*In-text superscripts \[A1\][^\n]*\n?", re.M)

JOBS = {"NEE": ("Manuscript_NEE_maintext", "Supplementary_Information_NEE", "numeric"),
        "NC":  ("Manuscript_NC_maintext",  "Supplementary_Information_NC",  "numeric"),
        "GCB": ("Manuscript_GCB_maintext", "Supporting_Information_GCB",    "authordate")}


def log(*a: object) -> None:
    print("[refs]", *a, flush=True)


def expand(tok: str) -> list[int]:
    out = []
    for p in tok.split(","):
        p = p.strip()
        if re.fullmatch(r"\d+", p):
            out.append(int(p))
        elif re.fullmatch(r"\d+[–-]\d+", p):
            a, b = re.split(r"[–-]", p)
            out.extend(range(int(a), int(b) + 1))
    return out


def split_main(text: str) -> tuple[str, str, str]:
    """正文 → (References 之前, 文献表, References 之后)。"""
    head, sep, rest = text.partition("## References")
    assert sep, "未找到 References 标题"
    m = re.search(r"\n## ", rest)
    return head, (rest[: m.start()] if m else rest), (rest[m.start():] if m else "")


def force_fonts(doc: Document, name: str = BODY_FONT) -> None:
    """把文档默认与各样式的主题字体换成显式字体 / replace theme fonts with an explicit face.

    pandoc 的样式(标题、题名、行号继承的 docDefaults)引用主题字体 Aptos;Word 中主题属性
    优先于显式 ascii 字体, 导致题名与标题并非 Times New Roman。代码样式保留等宽字体。
    NB: pandoc styles reference the theme font (Aptos), which Word prefers over an explicit
    face; strip the theme attributes everywhere except code styles.
    """
    root = doc.styles.element
    targets = [root.find(qn("w:docDefaults"))]
    for st in root.findall(qn("w:style")):
        nm = st.find(qn("w:name"))
        label = nm.get(qn("w:val")) if nm is not None else ""
        if not re.search(r"(?i)verbatim|source code|code", label):
            targets.append(st)
    for el in targets:
        if el is None:
            continue
        for rf in el.iter(qn("w:rFonts")):
            for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rf.attrib.pop(qn(att), None)
            for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(att), name)
    dd = root.find(qn("w:docDefaults"))                       # docDefaults 无 rFonts 时补上 / add if absent
    if dd is not None and not list(dd.iter(qn("w:rFonts"))):
        rpr = dd.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if rpr is not None:
            rf = OxmlElement("w:rFonts")
            for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(att), name)
            rpr.insert(0, rf)


_AFTER_SLN = {"pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
              "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing",
              "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
              "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"}


def _suppress_line_numbers(par) -> None:
    """按 OOXML 顺序插入 w:suppressLineNumbers / insert suppressLineNumbers in schema order."""
    ppr = par._p.get_or_add_pPr()
    if ppr.find(qn("w:suppressLineNumbers")) is not None:
        return
    el = OxmlElement("w:suppressLineNumbers")
    nxt = next((c for c in ppr if c.tag.split("}")[-1] in _AFTER_SLN), None)
    if nxt is not None:
        nxt.addprevious(el)
    else:
        ppr.append(el)


def style_document(path: Path, line_spacing: float | None = None) -> None:
    """稿件版式; line_spacing=None 为双倍行距(正文), SI 传 1.15 / manuscript layout; SI uses 1.15."""
    doc = Document(str(path))
    force_fonts(doc)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    if line_spacing is None:
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    else:                       # SI: 图与图注须同页(Nature 系要求)/ figure + legend on one page
        normal.paragraph_format.line_spacing = line_spacing
    normal.paragraph_format.space_after = Pt(0)
    for st in doc.styles:
        if st.name.startswith("Heading"):
            try:
                st.font.name, st.font.size, st.font.bold = BODY_FONT, Pt(BODY_PT + 1), True
                st.font.color.rgb = RGBColor(0, 0, 0)
            except AttributeError:
                pass
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(1)
        ln = OxmlElement("w:lnNumType")
        for k, v in (("countBy", "1"), ("start", "1"), ("restart", "continuous"), ("distance", "360")):
            ln.set(qn(f"w:{k}"), v)
        section._sectPr.append(ln)
        p = section.footer.paragraphs[0]
        p.alignment = 1
        # 页脚不参与行号(LibreOffice 会给页脚行编号, 导出的 PDF 出现第二个页码)
        # NB: keep the footer out of line numbering; LibreOffice otherwise numbers it in PDFs
        _suppress_line_numbers(p)
        for kind in ("begin", "instrText", "separate", "end"):
            el = OxmlElement("w:instrText" if kind == "instrText" else "w:fldChar")
            if kind == "instrText":
                el.set(qn("xml:space"), "preserve"); el.text = " PAGE "
            else:
                el.set(qn("w:fldCharType"), kind)
            p.add_run()._r.append(el)
    for tbl in doc.tables:
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "auto")
            borders.append(e)
        tbl._tbl.tblPr.append(borders)
        for row in tbl.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    para.paragraph_format.space_after = Pt(0)
                    for run in para.runs:
                        run.font.size, run.font.name = Pt(8), BODY_FONT
    doc.save(str(path))


def render(md_path: Path, docx_path: Path, line_spacing: float | None = None) -> None:
    r = subprocess.run(["pandoc", str(md_path), "-o", str(docx_path),
                        "--from", "markdown+raw_html+pipe_tables+implicit_figures",
                        "--to", "docx", "--standalone", "--resource-path", str(V2)],
                       capture_output=True, text=True)
    if r.returncode != 0:
        sys.exit(f"pandoc 失败 ({docx_path.name}): {r.stderr[:300]}")
    style_document(docx_path, line_spacing)


def first_surname_year(entry: str) -> tuple[str, str]:
    sur = re.match(r"([A-Z][A-Za-zÀ-ɏ'\-]+)", entry)
    yrs = re.findall(r"\((\d{4})[a-z]?\)", entry) or re.findall(r"\b(19\d{2}|20\d{2})\b", entry)
    return (sur.group(1) if sur else "?", yrs[-1] if yrs else "?")


def cited_keys(text: str) -> set[tuple[str, str]]:
    keys = set()
    for c in re.findall(r"\(([^()]*?(?:19|20)\d{2}[a-z]?[^()]*?)\)", text):
        for sur, yr in re.findall(
                r"([A-Z][A-Za-zÀ-ɏ'\-]+)(?:\s+(?:et al\.|&\s*[A-Z][A-Za-z'\-]+))?,?\s*((?:19|20)\d{2})", c):
            keys.add((sur, yr))
    return keys


def main() -> None:
    """裁剪、重编号并渲染三版 / prune, renumber and render all three versions."""
    for J, (main_stem, si_stem, mode) in JOBS.items():
        d = SUB / J
        main_md, si_md = d / f"{main_stem}.md", d / f"{si_stem}.md"
        main_txt, si_txt = main_md.read_text(encoding="utf-8"), si_md.read_text(encoding="utf-8")
        head, reflist, tail = split_main(main_txt)

        if mode == "numeric":
            entries = [l.strip() for l in reflist.split("\n") if re.match(r"^\d+\.\s", l.strip())]
            n0 = len(entries)
            # 作者/单位区的 ^1^ 是单位上标而非引文: 只在 Abstract 之后编号, 否则文献 1 被锁定在首位
            # NB: the affiliation ^1^ above the Abstract is not a citation; renumber only below it,
            #     otherwise reference 1 is pinned first regardless of where it is first cited.
            front, sep_ab, body = head.partition("## Abstract")
            if not sep_ab:
                front, body = "", head
            head = body
            # 首次出现顺序:正文(head) → References 之后的 Methods(tail) → SI
            order: list[int] = []
            for chunk in (head, tail, si_txt):
                for tok in SUP.findall(chunk):
                    for n in expand(tok):
                        if n not in order:
                            order.append(n)
            order = [n for n in order if 1 <= n <= n0]
            remap = {old: new for new, old in enumerate(order, 1)}

            def renum(txt: str) -> str:
                def fix(m: re.Match) -> str:
                    ns = [remap[n] for n in expand(m.group(1)) if n in remap]
                    if not ns:
                        return ""
                    ns = sorted(set(ns))
                    # 连续段折叠为 a–b / collapse runs
                    out, i = [], 0
                    while i < len(ns):
                        j = i
                        while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
                            j += 1
                        out.append(str(ns[i]) if j == i else f"{ns[i]}–{ns[j]}")
                        i = j + 1
                    return "^" + ",".join(out) + "^"
                return SUP.sub(fix, txt)

            kept = [re.sub(r"^\d+\.\s*", "", entries[o - 1]) for o in order]
            new_reflist = "\n\n".join(f"{i}. {e}" for i, e in enumerate(kept, 1))
            new_main = front + sep_ab + renum(head) + "## References\n\n" + new_reflist + "\n" + renum(tail)
            new_si = renum(si_txt)
        else:
            entries = [b.strip() for b in reflist.split("\n\n") if b.strip()]
            n0 = len(entries)
            keys = cited_keys(head) | cited_keys(tail) | cited_keys(si_txt)
            # 数据集/代码引用(机构作者难以按姓氏匹配)一律保留 —— 期刊要求在文献表中正式引用数据
            # NB: always keep data/code citations; corporate authors defeat surname matching.
            DATA_CIT = re.compile(r"Zenodo|datazone\.birdlife|GitHub")
            kept = [e for e in entries if first_surname_year(e) in keys or DATA_CIT.search(e)]
            new_reflist = "\n\n".join(kept)
            new_main = head + "## References\n\n" + new_reflist + "\n" + tail
            new_si = si_txt

        new_main = STALE.sub("", new_main)
        new_si = STALE.sub("", new_si)
        main_md.write_text(new_main, encoding="utf-8")
        si_md.write_text(new_si, encoding="utf-8")
        render(main_md, d / f"{main_stem}.docx")
        render(si_md, d / f"{si_stem}.docx", line_spacing=1.15)
        log(f"{J}: 文献 {n0} → {len(kept)} 条(裁掉 {n0-len(kept)});"
            f"{' 已按首次出现重编号;' if mode=='numeric' else ' 作者-年份制不重编号;'} 已重渲染 docx")


if __name__ == "__main__":
    main()
