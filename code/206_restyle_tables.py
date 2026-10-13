#!/usr/bin/env python3.11
# ============================================================
# 把投稿包内全部 Word 表格改为学术期刊通行的 booktabs 式。
# Restyle every table in the submission package to journal (booktabs) style.
#
# 问题 / Problem:
#   建包脚本给每张表加了完整网格线(上下左右 + 内部横竖线)。
#   期刊表格的通行体例是:只保留三条横线 —— 表顶、表头下、表底,
#   不使用任何竖线。满格网格是投稿表格最容易被一眼看出的非专业信号。
#
# 做法 / Approach:
#   - 表级边框:top / bottom 为实线(较粗),left / right / insideH / insideV 全部取消;
#   - 表头行:逐单元格加下边框(细),形成第二条横线;
#   - 单元格加左右内边距,避免文字贴线;
#   - 表头加粗、单元格垂直居中。
#   该脚本幂等,可对任意已生成的 .docx 重复运行。
#
# Input/Output: 原地重写 analysis_v2/submission/**/*.docx(presentation/ 与 work/ 除外)
# Run: python3.11 code/206_restyle_tables.py   (在 202 → 203/204 → 205 之后)
# ============================================================
from __future__ import annotations
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

SUB = Path("/Users/dingchenchen/Documents/New records/bird-new-distribution-records/"
           "tasks/bird_hazard_model_effort_upgrade_v2/analysis_v2/submission")


def _edge(tag: str, val: str, sz: int = 4) -> OxmlElement:
    e = OxmlElement(f"w:{tag}")
    e.set(qn("w:val"), val)
    if val != "none":
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), "auto")
    else:
        e.set(qn("w:sz"), "0")
        e.set(qn("w:color"), "auto")
    return e


def booktabs(tbl) -> None:
    """三线表:表顶、表头下、表底;无竖线。"""
    pr = tbl._tbl.tblPr
    for old in pr.findall(qn("w:tblBorders")):
        pr.remove(old)
    b = OxmlElement("w:tblBorders")
    b.append(_edge("top", "single", 8))
    b.append(_edge("bottom", "single", 8))
    for tag in ("left", "right", "insideH", "insideV"):
        b.append(_edge(tag, "none"))
    pr.append(b)
    # 单元格内边距,避免文字贴线 / cell padding
    for old in pr.findall(qn("w:tblCellMar")):
        pr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for side, w in (("top", 40), ("bottom", 40), ("left", 80), ("right", 80)):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), str(w)); m.set(qn("w:type"), "dxa")
        mar.append(m)
    pr.append(mar)

    for ci, cell in enumerate(tbl.rows[0].cells):          # 表头下的第二条横线
        tcPr = cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn("w:tcBorders")):
            tcPr.remove(old)
        tb = OxmlElement("w:tcBorders")
        tb.append(_edge("bottom", "single", 4))
        tcPr.append(tb)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True

    for row in tbl.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


if __name__ == "__main__":
    n_files = n_tables = 0
    # 只处理稿件与 SI;presentation/ 与 work/ 下的材料(如中英对照讲稿)不是投稿表格, 不得改动
    # NB: manuscripts and SI only - presentation/ and work/ hold non-submission documents.
    for f in sorted(SUB.rglob("*.docx")):
        if {"presentation", "work"} & set(f.relative_to(SUB).parts):
            continue
        doc = Document(str(f))
        if not doc.tables:
            continue
        for t in doc.tables:
            booktabs(t); n_tables += 1
        doc.save(str(f)); n_files += 1
        print(f"[tables] {f.relative_to(SUB)}: {len(doc.tables)} 张表已改为三线表")
    print(f"[tables] 完成:{n_files} 个文件,{n_tables} 张表")
